
import os
import io
import re
import traceback
from datetime import datetime

from flask import (
    Flask, request, render_template_string,
    send_file, flash, redirect, url_for
)
from werkzeug.utils import secure_filename

from reportlab.pdfgen import canvas
from reportlab.lib.units import inch
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase.pdfmetrics import stringWidth

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None


# ==================================================
# SETTINGS
# ==================================================

app = Flask(__name__)
app.secret_key = "stanns-question-paper-generator"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
HEADER_IMAGE = os.path.join(BASE_DIR, "college_header.jpg")
UPLOAD_FOLDER = os.path.join(BASE_DIR, "uploads")

os.makedirs(UPLOAD_FOLDER, exist_ok=True)

# CUSTOM SIZE - NOT A4
PAGE_WIDTH = 5.8 * inch
PAGE_HEIGHT = 8.26 * inch
PAGE_SIZE = (PAGE_WIDTH, PAGE_HEIGHT)

CURRENT_PAPER = {
    "topic": "",
    "exam_time": "3 Hours",
    "max_marks": "60",
    "section_a_marks": "30",
    "section_b_marks": "30",
    "questions_a": [],
    "questions_b": [],
}


# ==================================================
# WEBSITE HTML
# ==================================================

HTML_PAGE = r"""
<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">

<title>St. Ann's Question Paper Generator</title>

<style>
* { box-sizing: border-box; }

body {
    margin: 0;
    padding: 18px;
    background: #edf2f7;
    font-family: Arial, sans-serif;
    color: #222;
}

.container {
    max-width: 1050px;
    margin: auto;
}

h1 {
    text-align: center;
    color: #007363;
}

.form-panel {
    background: white;
    padding: 20px;
    margin: 20px 0;
    border-radius: 8px;
    box-shadow: 0 2px 10px #00000012;
}

label {
    display: block;
    font-weight: bold;
    margin: 12px 0 6px;
}

input[type="text"],
input[type="number"],
input[type="file"] {
    width: 100%;
    padding: 10px;
    border: 1px solid #b8c2cc;
    border-radius: 5px;
    background: white;
}

.two-columns {
    display: grid;
    grid-template-columns: 1fr 1fr;
    gap: 15px;
}

button, .download {
    display: inline-block;
    padding: 12px 17px;
    margin: 10px 5px 0 0;
    background: #007363;
    color: white;
    border: 0;
    border-radius: 5px;
    cursor: pointer;
    font-size: 15px;
    text-decoration: none;
}

button:hover, .download:hover {
    background: #00574c;
}

.message {
    padding: 12px;
    margin: 8px 0;
    background: #fff3cd;
    border: 1px solid #e5d49b;
    border-radius: 5px;
}

.preview-title {
    text-align: center;
    color: #007363;
}

.preview-scroll {
    width: 100%;
    overflow-x: auto;
    padding: 12px 0 20px;
}

.paper {
    width: 5.8in;
    min-width: 5.8in;
    min-height: 8.26in;
    padding: 0.28in;
    margin: 0 auto;
    background: white;
    color: black;
    border: 1px solid #bbb;
    font-family: "Times New Roman", serif;
    font-size: 10pt;
    overflow-wrap: anywhere;
}

.letterhead {
    display: block;
    width: 100%;
    height: auto;
    margin: 0 auto 10px;
}

.college-name {
    text-align: center;
    font-weight: bold;
    font-size: 12pt;
    line-height: 1.3;
    margin-bottom: 12px;
}

.topic-title {
    text-align: center;
    font-weight: bold;
    text-decoration: underline;
    margin: 10px 0;
}

.paper-info {
    display: flex;
    justify-content: space-between;
    gap: 10px;
    font-weight: bold;
    margin: 14px 0;
}

.section-title {
    font-weight: bold;
    margin: 12px 0 7px;
}

.question {
    margin: 6px 0;
    line-height: 1.3;
}

@media (max-width: 600px) {
    body { padding: 8px; }
    .form-panel { padding: 12px; }
    .two-columns { grid-template-columns: 1fr; }
    .paper { margin-left: 0; }
}
</style>
</head>

<body>
<div class="container">

<h1>St. Ann's Question Paper Generator</h1>

<div class="form-panel">
    {% for message in get_flashed_messages() %}
        <div class="message">{{ message }}</div>
    {% endfor %}

    <form method="POST" enctype="multipart/form-data">

        <label>Upload Lesson Plan (PDF or Word DOCX)</label>
        <input type="file" name="lesson_file"
               accept=".pdf,.docx" required>

        <label>Unit Topic / Topic Heading</label>
        <input type="text" name="topic"
               value="{{ topic }}"
               placeholder="Enter unit topic or heading">

        <div class="two-columns">
            <div>
                <label>Time</label>
                <input type="text" name="exam_time"
                       value="{{ exam_time }}"
                       placeholder="Example: 3 Hours">
            </div>

            <div>
                <label>Maximum Marks</label>
                <input type="number" name="max_marks"
                       value="{{ max_marks }}" min="1" required>
            </div>
        </div>

        <div class="two-columns">
            <div>
                <label>Section A Total Marks</label>
                <input type="number" name="section_a_marks"
                       value="{{ section_a_marks }}" min="0">
            </div>

            <div>
                <label>Section B Total Marks</label>
                <input type="number" name="section_b_marks"
                       value="{{ section_b_marks }}" min="0">
            </div>
        </div>

        <button type="submit">Generate Question Paper</button>
    </form>
</div>

{% if questions_a and questions_b %}
<h2 class="preview-title">
    Preview — Custom Page Size 5.8 × 8.26 inches
</h2>

<div class="preview-scroll">
<div class="paper">

    {% if header_exists %}
        <img class="letterhead"
             src="{{ url_for('header_image') }}"
             alt="College Letterhead">
    {% else %}
        <div class="college-name">
            St. Ann's College for Women (A), Autonomous,<br>
            NAAC 'A' Grade
        </div>
    {% endif %}

    <div class="topic-title">
        {{ topic or "QUESTION PAPER" }}
    </div>

    <div class="paper-info">
        <span>Time: {{ exam_time or "________" }}</span>
        <span>Maximum Marks: {{ max_marks }}</span>
    </div>

    <div class="section-title">
        SECTION A
        {% if section_a_marks %} ({{ section_a_marks }} Marks){% endif %}
    </div>

    {% for question in questions_a %}
        <div class="question">{{ loop.index }}. {{ question }}</div>
    {% endfor %}

    <div class="section-title">
        SECTION B
        {% if section_b_marks %} ({{ section_b_marks }} Marks){% endif %}
    </div>

    {% for question in questions_b %}
        <div class="question">{{ loop.index }}. {{ question }}</div>
    {% endfor %}

</div>
</div>

<div class="form-panel">
    <a class="download"
       href="{{ url_for('download_file', kind='pdf') }}">
       Download PDF
    </a>

    <a class="download"
       href="{{ url_for('download_file', kind='word') }}">
       Download Word
    </a>
</div>
{% endif %}

</div>
</body>
</html>
"""


# ==================================================
# READ UPLOADED DOCUMENT
# ==================================================

def allowed_file(filename):
    return (
        "." in filename
        and filename.rsplit(".", 1)[1].lower() in {"pdf", "docx"}
    )


def read_lesson_file(path):
    extension = os.path.splitext(path)[1].lower()

    if extension == ".pdf":
        if PdfReader is None:
            raise ValueError(
                "pypdf is not installed. Run: python -m pip install pypdf"
            )

        reader = PdfReader(path)
        parts = []

        for page in reader.pages:
            parts.append(page.extract_text() or "")

        text = "\n".join(parts)

    elif extension == ".docx":
        document = Document(path)
        parts = [p.text for p in document.paragraphs if p.text.strip()]

        for table in document.tables:
            for row in table.rows:
                parts.append(
                    " ".join(cell.text for cell in row.cells)
                )

        text = "\n".join(parts)

    else:
        raise ValueError("Upload a PDF or Word DOCX file.")

    text = re.sub(r"\x00", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = text.strip()

    if not text:
        raise ValueError(
            "PDFలో చదవగలిగే text లేదు. "
            "ఇది scanned PDF అయితే OCR చేసిన PDF లేదా DOCX upload చేయండి."
        )

    return text


# ==================================================
# GENERATE QUESTIONS FROM LESSON TEXT
# ==================================================

def make_questions(lesson_text, supplied_topic):
    lines = []

    for line in lesson_text.splitlines():
        line = re.sub(r"\s+", " ", line).strip()
        line = re.sub(r"^[\s\d.\-*)]+", "", line).strip()

        if len(line) >= 8:
            lines.append(line)

    # Remove duplicate lines.
    unique_lines = []
    seen = set()

    for line in lines:
        if line.lower() not in seen:
            unique_lines.append(line)
            seen.add(line.lower())

    if not unique_lines:
        raise ValueError(
            "Lesson documentలో ఉపయోగించగల text కనిపించలేదు."
        )

    topic = supplied_topic.strip() if supplied_topic.strip() else unique_lines[0]

    # Use lesson lines as the question concepts.
    concepts = []

    for line in unique_lines:
        if re.search(
            r"college name|signature|page \d+|attendance",
            line,
            re.IGNORECASE
        ):
            continue

        # Keep long paragraphs from making excessively long questions.
        if len(line) > 95:
            line = line[:92].rsplit(" ", 1)[0] + "..."

        concepts.append(line)

    if not concepts:
        concepts = [topic]

    section_a_templates = [
        "Define {x}.",
        "What is meant by {x}?",
        "Write a short note on {x}.",
        "State the importance of {x}.",
        "Mention the main features of {x}.",
        "Explain the meaning of {x}.",
        "List the important points related to {x}.",
        "What are the objectives of {x}?",
        "Identify the key aspects of {x}.",
        "Briefly describe {x}."
    ]

    section_b_templates = [
        "Explain {x} in detail.",
        "Discuss the importance and features of {x}.",
        "Describe the main concepts related to {x}.",
        "Explain the objectives and applications of {x}.",
        "Discuss the advantages and limitations of {x}.",
        "Describe the principles of {x}.",
        "Explain the major components of {x}.",
        "Write a detailed note on {x}.",
        "Analyse the role and significance of {x}.",
        "Explain {x} with suitable examples."
    ]

    questions_a = []
    questions_b = []

    for i in range(10):
        concept = concepts[i % len(concepts)]
        questions_a.append(
            section_a_templates[i].format(x=concept)
        )

    for i in range(10):
        concept = concepts[(i + 10) % len(concepts)]
        questions_b.append(
            section_b_templates[i].format(x=concept)
        )

    return topic, questions_a, questions_b


# ==================================================
# PDF GENERATOR - CUSTOM SIZE
# ==================================================

def wrap_text(text, font_name, font_size, max_width):
    words = str(text).split()
    result = []
    current = ""

    for word in words:
        trial = word if not current else current + " " + word

        if stringWidth(trial, font_name, font_size) <= max_width:
            current = trial
        else:
            if current:
                result.append(current)
            current = word

    if current:
        result.append(current)

    return result or [""]


def create_pdf(data):
    output = io.BytesIO()
    pdf = canvas.Canvas(output, pagesize=PAGE_SIZE)

    pdf.setTitle("St. Ann's Question Paper")
    pdf.setAuthor("St. Ann's College for Women")

    margin = 0.32 * inch
    usable_width = PAGE_WIDTH - 2 * margin
    y = PAGE_HEIGHT - 0.30 * inch

    def start_page():
        nonlocal y

        y = PAGE_HEIGHT - 0.30 * inch

        if os.path.isfile(HEADER_IMAGE):
            try:
                image = ImageReader(HEADER_IMAGE)
                image_width, image_height = image.getSize()

                max_width = usable_width
                max_height = 1.30 * inch
                scale = min(
                    max_width / image_width,
                    max_height / image_height
                )

                draw_width = image_width * scale
                draw_height = image_height * scale

                pdf.drawImage(
                    image,
                    (PAGE_WIDTH - draw_width) / 2,
                    y - draw_height,
                    width=draw_width,
                    height=draw_height,
                    preserveAspectRatio=True,
                    mask="auto"
                )

                y -= draw_height + 0.10 * inch

            except Exception:
                pass
        else:
            pdf.setFont("Times-Bold", 10)
            pdf.drawCentredString(
                PAGE_WIDTH / 2, y - 0.15 * inch,
                "St. Ann's College for Women (A), Autonomous,"
            )
            pdf.drawCentredString(
                PAGE_WIDTH / 2, y - 0.34 * inch,
                "NAAC 'A' Grade"
            )
            y -= 0.52 * inch

    def write_lines(text, font="Times-Roman", size=8.5, indent=0):
        nonlocal y

        lines = wrap_text(
            text, font, size, usable_width - indent
        )
        leading = size * 1.25

        for line in lines:
            if y - leading < 0.35 * inch:
                pdf.showPage()
                start_page()

            pdf.setFont(font, size)
            pdf.drawString(margin + indent, y, line)
            y -= leading

    start_page()

    write_lines(
        data["topic"] or "QUESTION PAPER",
        "Times-Bold", 10
    )
    y -= 4

    if y < 0.60 * inch:
        pdf.showPage()
        start_page()

    pdf.setFont("Times-Bold", 8.5)
    pdf.drawString(
        margin, y,
        "Time: " + (data["exam_time"] or "________")
    )
    pdf.drawRightString(
        PAGE_WIDTH - margin, y,
        "Maximum Marks: " + str(data["max_marks"])
    )
    y -= 0.25 * inch

    write_lines(
        "SECTION A (" + str(data["section_a_marks"]) + " Marks)",
        "Times-Bold", 9.5
    )
    y -= 2

    for index, question in enumerate(data["questions_a"], 1):
        write_lines(
            str(index) + ". " + question,
            "Times-Roman", 8, 5
        )
        y -= 2

    y -= 3

    write_lines(
        "SECTION B (" + str(data["section_b_marks"]) + " Marks)",
        "Times-Bold", 9.5
    )
    y -= 2

    for index, question in enumerate(data["questions_b"], 1):
        write_lines(
            str(index) + ". " + question,
            "Times-Roman", 8, 5
        )
        y -= 2

    pdf.save()
    output.seek(0)
    return output


# ==================================================
# WORD GENERATOR - CUSTOM SIZE
# ==================================================

def create_word(data):
    output = io.BytesIO()
    document = Document()

    section = document.sections[0]
    section.page_width = Inches(5.8)
    section.page_height = Inches(8.26)
    section.top_margin = Inches(0.28)
    section.bottom_margin = Inches(0.30)
    section.left_margin = Inches(0.32)
    section.right_margin = Inches(0.32)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(9)

    if os.path.isfile(HEADER_IMAGE):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(4)

        try:
            paragraph.add_run().add_picture(
                HEADER_IMAGE, width=Inches(5.0)
            )
        except Exception:
            paragraph.add_run(
                "St. Ann's College for Women (A), Autonomous,\n"
                "NAAC 'A' Grade"
            ).bold = True
    else:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(
            "St. Ann's College for Women (A), Autonomous,\n"
            "NAAC 'A' Grade"
        )
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(5)

    run = paragraph.add_run(data["topic"] or "QUESTION PAPER")
    run.bold = True
    run.underline = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)

    # Time on left, maximum marks on right.
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(2.45)
    table.columns[1].width = Inches(2.45)

    left = table.cell(0, 0).paragraphs[0]
    left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    left.add_run(
        "Time: " + (data["exam_time"] or "________")
    ).bold = True

    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.add_run(
        "Maximum Marks: " + str(data["max_marks"])
    ).bold = True

    def add_section(title, marks, questions):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(7)
        paragraph.paragraph_format.space_after = Pt(3)

        run = paragraph.add_run(
            title + " (" + str(marks) + " Marks)"
        )
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(9.5)

        for index, question in enumerate(questions, 1):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.05)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.keep_together = True

            run = paragraph.add_run(
                str(index) + ". " + question
            )
            run.font.name = "Times New Roman"
            run.font.size = Pt(8)

    add_section(
        "SECTION A", data["section_a_marks"], data["questions_a"]
    )
    add_section(
        "SECTION B", data["section_b_marks"], data["questions_b"]
    )

    document.save(output)
    output.seek(0)
    return output


# ==================================================
# WEBSITE ROUTES
# ==================================================

@app.route("/college_header.jpg")
def header_image():
    if not os.path.isfile(HEADER_IMAGE):
        return "college_header.jpg not found", 404

    return send_file(HEADER_IMAGE, mimetype="image/jpeg")


@app.route("/", methods=["GET", "POST"])
def index():
    if request.method == "POST":
        file_path = None

        try:
            uploaded = request.files.get("lesson_file")

            topic_input = request.form.get("topic", "").strip()
            exam_time = request.form.get("exam_time", "").strip()
            max_marks = request.form.get("max_marks", "60").strip()
            section_a_marks = request.form.get(
                "section_a_marks", "30"
            ).strip()
            section_b_marks = request.form.get(
                "section_b_marks", "30"
            ).strip()

            if not uploaded or not uploaded.filename:
                flash("దయచేసి PDF లేదా Word Lesson Plan ఎంచుకోండి.")
                return redirect(url_for("index"))

            if not allowed_file(uploaded.filename):
                flash("PDF లేదా Word DOCX ఫైల్ మాత్రమే upload చేయండి.")
                return redirect(url_for("index"))

            try:
                if int(max_marks) < 1:
                    raise ValueError
            except (ValueError, TypeError):
                flash("Maximum Marksలో సరైన సంఖ్య నమోదు చేయండి.")
                return redirect(url_for("index"))

            filename = secure_filename(uploaded.filename)
            unique_name = (
                datetime.now().strftime("%Y%m%d%H%M%S%f")
                + "_" + filename
            )
            file_path = os.path.join(UPLOAD_FOLDER, unique_name)

            uploaded.save(file_path)

            lesson_text = read_lesson_file(file_path)

            topic, questions_a, questions_b = make_questions(
                lesson_text, topic_input
            )

            CURRENT_PAPER.update({
                "topic": topic,
                "exam_time": exam_time,
                "max_marks": max_marks,
                "section_a_marks": section_a_marks,
                "section_b_marks": section_b_marks,
                "questions_a": questions_a,
                "questions_b": questions_b,
            })

            flash("Question Paper తయారైంది. Previewను చూడండి.")

        except Exception as exc:
            print("\nERROR WHILE GENERATING QUESTION PAPER:")
            traceback.print_exc()
            flash("Question Paper తయారు కాలేదు: " + str(exc))

        finally:
            if file_path and os.path.isfile(file_path):
                try:
                    os.remove(file_path)
                except OSError:
                    pass

    return render_template_string(
        HTML_PAGE,
        topic=CURRENT_PAPER["topic"],
        exam_time=CURRENT_PAPER["exam_time"],
        max_marks=CURRENT_PAPER["max_marks"],
        section_a_marks=CURRENT_PAPER["section_a_marks"],
        section_b_marks=CURRENT_PAPER["section_b_marks"],
        questions_a=CURRENT_PAPER["questions_a"],
        questions_b=CURRENT_PAPER["questions_b"],
        header_exists=os.path.isfile(HEADER_IMAGE),
    )


@app.route("/download/<kind>")
def download_file(kind):
    data = CURRENT_PAPER

    if not data["questions_a"] or not data["questions_b"]:
        flash("ముందుగా Lesson Plan upload చేసి Generate క్లిక్ చేయండి.")
        return redirect(url_for("index"))

    try:
        if kind == "pdf":
            output = create_pdf(data)

            return send_file(
                output,
                as_attachment=True,
                download_name="St_Anns_Question_Paper.pdf",
                mimetype="application/pdf",
            )

        if kind == "word":
            output = create_word(data)

            return send_file(
                output,
                as_attachment=True,
                download_name="St_Anns_Question_Paper.docx",
                mimetype=(
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
            )

        flash("తెలియని download format.")
        return redirect(url_for("index"))

    except Exception as exc:
        print("\nERROR WHILE DOWNLOADING:")
        traceback.print_exc()
        flash("Download సమస్య: " + str(exc))
        return redirect(url_for("index"))


# ==================================================
# START APPLICATION
# ==================================================

if __name__ == "__main__":
    print("St. Ann's Question Paper Generator")
    print("Open: http://127.0.0.1:5000")
    app.run(host="0.0.0.0", port=5000, debug=False)
